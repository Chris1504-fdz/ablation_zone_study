"""The Campaign class: the BO campaign log and the five iteration steps."""
import os
import numpy as np
import pandas as pd
from IPython.display import display, Markdown
from .config import SEED, X_COLS, Y_COLS, X_NAMES, Y_NAMES, Y_SHORT, I369_WT
from .data import condition_stats
from .surrogate import DesignSpace, fit_surrogate, predict, propose, acquisition_map, fit_quality, gp_hyperparameters
from . import plots
from . import stopping as prb


class Campaign:
    """Holds the master dataset, the iteration tags, the design space and one record per
    iteration. In every iteration section of the notebook the steps are called in order:

        propose(it, pinned=None, sent=None, note='')   sample requested from UIC (+ predictions)
        import_results(it, observed)                   replicates provided by UIC
        compare(it)                                    provided (predicted) vs. received (observed)
        update(it)                                     Pareto front & hypervolume through `it`
        improvement(it)                                best-so-far per output, HV growth, per-sample gains

    Stopping rule (Wilson, NeurIPS 2024): stopping_check(it) before requesting iteration `it`,
    stopping_trajectory() for how the rule evolved over the campaign.
    Summaries: baseline(), table(), prediction_accuracy(), response_surfaces(), export(), summary().
    """

    def __init__(self, df_rep, iteration_samples, out_dir='../outputs/v4', data_file='', seed=SEED):
        self.samples = {int(k): [tuple(map(float, xy)) for xy in v] for k, v in iteration_samples.items()}
        self.records, self.stopping = {}, {}
        self.seed, self.out_dir, self.data_file = seed, out_dir, data_file
        os.makedirs(out_dir, exist_ok=True)
        self.df_full = df_rep[X_COLS + Y_COLS].astype(float).reset_index(drop=True).copy()
        self.df_full.insert(0, 'Iteration', [self.iteration_of(r.X1, r.X2) for r in self.df_full.itertuples()])
        self.stats0, _, _ = self.condition_stats(self.replicates_through(0))
        self.space = DesignSpace(self.stats0)

    # ────────────────────────────── data access ──────────────────────────────
    def iteration_of(self, x1, x2):
        """Campaign iteration a condition belongs to (0 = initial DOE). Also recognises
        results imported with `import_results` that are not in the spreadsheet yet."""
        for it, conds in self.samples.items():
            if any(np.isclose(x1, a) and np.isclose(x2, b) for a, b in conds):
                return it
        for it, rec in self.records.items():
            if 'observed' in rec and np.allclose(rec['x_fab'][0], [x1, x2]):
                return it
        return 0

    def condition_stats(self, df_rep):
        return condition_stats(df_rep, self.iteration_of)

    def replicates_through(self, it):
        """All replicate rows through iteration `it`. Imported results missing from the
        spreadsheet are appended automatically (and ignored once they are in it)."""
        out = self.df_full[self.df_full.Iteration <= it].drop(columns='Iteration')
        for k in range(1, it + 1):
            obs = self.records.get(k, {}).get('observed')
            if obs is None:
                continue
            for x1, x2 in obs[X_COLS].drop_duplicates().values:
                if not (np.isclose(out.X1, x1) & np.isclose(out.X2, x2)).any():
                    out = pd.concat([out, obs[np.isclose(obs.X1, x1) & np.isclose(obs.X2, x2)][X_COLS + Y_COLS]],
                                    ignore_index=True)
        return out.reset_index(drop=True)

    def pareto_summary(self, it):
        """Data through `it` → stats, GP noise, Pareto mask and hypervolume in both conventions."""
        df_rep = self.replicates_through(it)
        stats, Yvar, s2 = self.condition_stats(df_rep)
        pm, hv_means = self.space.pareto(stats[Y_COLS].values)
        doe = stats.loc[stats.Iteration == 0, Y_COLS].values
        reps = df_rep[[self.iteration_of(a, b) > 0 for a, b in df_rep[X_COLS].values]][Y_COLS].values
        pmr, hv_reps = self.space.pareto(np.vstack([doe, reps]) if len(reps) else doe)
        return dict(df_rep=df_rep, stats=stats, Yvar=Yvar, s2=s2, pm=pm, hv_means=hv_means,
                    n_pareto=int(pm.sum()), hv_reps=hv_reps, n_pareto_reps=int(pmr.sum()))

    @property
    def tested(self):
        return [it for it in sorted(self.records) if 'observed' in self.records[it]]

    @property
    def pending(self):
        return [it for it in sorted(self.records) if 'observed' not in self.records[it]]

    # ────────────────────────────── overview / baseline ──────────────────────
    def overview(self):
        """Section 1: what the spreadsheet contains, per iteration."""
        df = self.df_full
        display(Markdown(f"**{len(df)} measurements / {len(df[X_COLS].drop_duplicates())} conditions** "
                         f"loaded from `{self.data_file}`"))
        display(pd.DataFrame({
            'rows': df.groupby('Iteration').size(),
            'conditions': df.groupby('Iteration')[X_COLS].apply(lambda g: len(g.drop_duplicates())),
            'samples (X1, X2)': [', '.join(f'({a:g}, {b:g})' for a, b in self.samples.get(it, [])) or 'initial DOE'
                                 for it in sorted(df.Iteration.unique())]}))
        display(Markdown('**Campaign rows in the spreadsheet**'))
        display(df[df.Iteration > 0].reset_index(drop=True))
        display(Markdown('**Initial DOE — condition means, replicate variances s² and n** (GP noise = s²/n)'))
        display(self.stats0.drop(columns='Iteration').round(4))
        _, _, s2 = self.condition_stats(self.replicates_through(0))
        print(f'Pooled replicate variance s̄²:  Y1 {s2[0]:.4f}  |  Y2 {s2[1]:.2f}  |  Y3 {s2[2]:.1e}')
        print('Hypervolume reference point (worse than the DOE nadir by 10%): '
              + ' | '.join(f'{c} > {v:.4g}' for c, v in zip(Y_COLS, self.space.ref_point_orig())))

    def baseline(self):
        """Section 3: surrogate fit, Pareto front and hypervolume of the initial DOE."""
        ps = self.pareto_summary(0)
        model, _ = fit_surrogate(self.space, ps['stats'], ps['Yvar'])
        display(Markdown('**GP hyperparameters (initial fit)**')); display(gp_hyperparameters(model).round(4))
        display(Markdown('**In-sample fit quality**')); display(fit_quality(self.space, model, ps['stats']).round(4))
        print(f"Pareto front: {ps['n_pareto']} / {len(ps['stats'])} conditions  |  HV = {ps['hv_means']:.4f}")
        display(Markdown('**Pareto-optimal conditions of the initial DOE**'))
        display(ps['stats'].loc[ps['pm'], X_COLS + Y_COLS].round(4))
        plots.design_space(ps['stats'], title='Design space — initial DOE')
        plots.pareto_pairs(ps['stats'], ps['pm'],
                           f"Initial DOE — Pareto front ({ps['n_pareto']} solutions)  |  HV = {ps['hv_means']:.3f}")
        return ps

    # ────────────────────────────── the five steps ───────────────────────────
    def propose(self, it, pinned=None, sent=None, note=''):
        """Step 1 — PROPOSE. Fit on data through it−1, maximize qLogEHVI, and record the
        sample requested from UIC together with the model prediction there.
        pinned = {'x': (X1, X2), 'ehvi': …, 'origin': …} fixes a historical request (what was
        actually sent); the optimizer still runs and is reported for comparison.
        sent   = DataFrame(mean, low, high) of the predictions that accompanied the request."""
        ps = self.pareto_summary(it - 1)
        model, Yneg = fit_surrogate(self.space, ps['stats'], ps['Yvar'])
        x_run, ehvi_run = propose(self.space, model, Yneg, seed=self.seed)
        if pinned is None:
            x, ehvi, origin = x_run, ehvi_run, f'v4 optimizer (this run, seed {self.seed})'
        else:
            x = np.atleast_2d(np.asarray(pinned['x'], dtype=float))
            ehvi, origin = pinned['ehvi'], pinned.get('origin', 'campaign record')
        pred = predict(self.space, model, x, ps['s2'])
        rec = dict(it=it, x_proposed=x, x_fab=np.round(x, 1), ehvi=ehvi, origin=origin, note=note,
                   x_run=x_run, ehvi_run=ehvi_run, model=model, Yneg=Yneg, stats=ps['stats'],
                   s2_pooled=ps['s2'], pred=pred, sent=sent,
                   sent_source=(sent.attrs.get('source', '') if sent is not None else ''))
        self.records[it] = rec

        display(Markdown(f'#### Iteration {it} · Step 1 — sample requested from UIC'))
        display(pd.DataFrame([{
            'Trained on': f"{len(ps['stats'])} conditions (through iteration {it - 1})",
            'Requested X1': x[0, 0], 'Requested X2': x[0, 1],
            'Fabricated as': f"{rec['x_fab'][0, 0]:g} / {rec['x_fab'][0, 1]:g}",
            'Resin wt%': round(100 - rec['x_fab'][0, 0] - rec['x_fab'][0, 1] - I369_WT, 1),
            'qLogEHVI': ehvi, 'Origin': origin, 'Note': note}]).T.rename(columns={0: f'iteration {it}'}))
        msg = f'v4 optimizer on the same data: X1 = {x_run[0, 0]:.3f}, X2 = {x_run[0, 1]:.3f} (qLogEHVI {ehvi_run:.4f})'
        if pinned is not None:
            d = np.abs(x_run - x).max()
            msg += f" → {'matches' if d < 0.15 else 'differs from'} the requested sample (max |Δ| = {d:.2f})"
        print(msg)
        display(fit_quality(self.space, model, ps['stats']).round(4))
        self._show_prediction(pred, f'v4 model prediction at the requested sample (X1 = {x[0, 0]:.3f}, X2 = {x[0, 1]:.3f})')
        if sent is not None:
            display(Markdown(f"**Predictions that were sent with the request** — source: `{rec['sent_source']}`"))
            display(sent.round(4))
        X1g, X2g, acq = acquisition_map(self.space, model, Yneg)
        plots.acquisition(X1g, X2g, acq, ps['stats'], x,
                          f'Iteration {it} — qLogEHVI landscape (data through iteration {it - 1})',
                          'Requested sample' if pinned is not None else 'Proposed sample')
        return rec

    def import_results(self, it, observed):
        """Step 2 — IMPORT. Attach the replicates UIC provided for iteration `it` and
        cross-check them against the spreadsheet."""
        rec = self.records[it]
        rec['source'] = observed.attrs.get('source', 'n/a')
        rec['observed'] = observed.reset_index(drop=True)
        x1, x2 = float(observed.X1.iloc[0]), float(observed.X2.iloc[0])
        assert np.allclose([x1, x2], rec['x_fab'][0]), 'observed condition ≠ fabricated composition'
        display(Markdown(f"#### Iteration {it} · Step 2 — results provided by UIC  \nsource: `{rec['source']}`"))
        show = rec['observed'].copy(); show.insert(0, 'replicate', range(1, len(show) + 1))
        show.loc['mean'] = ['—', x1, x2] + [rec['observed'][c].mean() for c in Y_COLS]
        display(show)
        in_sheet = self.df_full[np.isclose(self.df_full.X1, x1) & np.isclose(self.df_full.X2, x2)]
        if len(in_sheet):
            same = np.allclose(np.sort(in_sheet[Y_COLS].values, axis=0), np.sort(rec['observed'][Y_COLS].values, axis=0))
            print(f"Spreadsheet check: ({x1:g}, {x2:g}) present with {len(in_sheet)} rows → "
                  f"{'✓ identical to the imported replicates' if same else '⚠ values differ from the imported replicates'}")
        else:
            print(f'Spreadsheet check: ({x1:g}, {x2:g}) not in the spreadsheet yet → the imported replicates are '
                  f'used from here on (append them to the spreadsheet and to ITERATION_SAMPLES when convenient)')

    def compare(self, it):
        """Step 3 — COMPARE. Observed replicates vs. the interval that was sent and vs. the
        v4 model's CI/PI; z = (ȳ − µ) / √(σ²_latent + s̄²/n) for the condition mean."""
        rec = self.records[it]
        obs, pred, sent = rec['observed'], rec['pred'], rec.get('sent')
        n, rows = len(obs), []
        for k, c in enumerate(Y_COLS):
            o = obs[c].values; p = pred.loc[k]
            r = {'Objective': Y_SHORT[c]}
            if sent is not None:
                s = sent.loc[k]
                r['Sent µ'] = s['mean']; r['Sent 95% range'] = f"[{s['low']:.4g}, {s['high']:.4g}]"
            r['v4 µ'] = p['mean']; r['v4 95% PI'] = f"[{p['pi_low']:.4g}, {p['pi_high']:.4g}]"
            r['Observed replicates'] = ', '.join(f'{v:g}' for v in o); r['Obs mean'] = o.mean()
            if sent is not None:
                r['Δ obs − sent'] = o.mean() - s['mean']
                r['In sent range'] = f"{int(((o >= s['low']) & (o <= s['high'])).sum())}/{n}"
            r['Δ obs − v4'] = o.mean() - p['mean']
            r['In v4 CI'] = f"{int(((o >= p['ci_low']) & (o <= p['ci_high'])).sum())}/{n}"
            r['In v4 PI'] = f"{int(((o >= p['pi_low']) & (o <= p['pi_high'])).sum())}/{n}"
            r['z (obs mean)'] = (o.mean() - p['mean']) / np.sqrt(p['sd_latent'] ** 2 + rec['s2_pooled'][k] / n)
            rows.append(r)
        table = pd.DataFrame(rows)
        display(Markdown(f'#### Iteration {it} · Step 3 — provided (predicted) vs. received (observed)'))
        display(table.round(4))
        plots.compare(it, obs, pred, sent)
        rec['comparison'] = table
        return table

    def update(self, it):
        """Step 4 — UPDATE. Merge iteration `it`, recompute the Pareto front and hypervolume,
        and say what the new sample changed."""
        ps, prev = self.pareto_summary(it), self.pareto_summary(it - 1)
        rec = self.records[it]
        rec.update(hv_means=ps['hv_means'], hv_reps=ps['hv_reps'], n_pareto=ps['n_pareto'])
        st = ps['stats']; new = (st.Iteration == it).values
        display(Markdown(f'#### Iteration {it} · Step 4 — updated dataset, Pareto front & hypervolume'))
        display(pd.DataFrame(
            [{'Conditions': len(prev['stats']), 'Pareto size': prev['n_pareto'], 'HV (means)': prev['hv_means'],
              'HV (reps)': prev['hv_reps'], 'Pareto size (reps)': prev['n_pareto_reps']},
             {'Conditions': len(st), 'Pareto size': ps['n_pareto'], 'HV (means)': ps['hv_means'],
              'HV (reps)': ps['hv_reps'], 'Pareto size (reps)': ps['n_pareto_reps']}],
            index=[f'through iteration {it - 1}', f'through iteration {it}']).round(4))
        print(f"HV gain this iteration: {ps['hv_means'] - prev['hv_means']:+.4f} on condition means "
              f"({100 * (ps['hv_means'] / prev['hv_means'] - 1):+.2f}%), "
              f"{ps['hv_reps'] - prev['hv_reps']:+.4f} with replicates as points")
        for c in Y_COLS:
            best_before, val = st.loc[~new, c].min(), st.loc[new, c].iloc[0]
            print(f"  • {Y_SHORT[c]}: new sample {val:.4g} vs best before {best_before:.4g} → "
                  f"{'NEW BEST' if val < best_before else 'not the best'}")
        y_new, y_prev = st.loc[new, Y_COLS].values[0], st.loc[~new, Y_COLS].values
        dominated = np.all(y_new <= y_prev, axis=1) & np.any(y_new < y_prev, axis=1)
        print(f"  • New sample {'is ON' if ps['pm'][new].all() else 'is NOT on'} the Pareto front; "
              f"it dominates {int(dominated.sum())} previous condition(s)")
        plots.pareto_pairs(st, ps['pm'], f"Pareto front through iteration {it} ({ps['n_pareto']} solutions)  |  "
                                         f"HV = {ps['hv_means']:.3f} (previous {prev['hv_means']:.3f})")
        return ps

    def improvement(self, through_it, full=False):
        """Step 5 — TRACK. Best condition mean so far per output (and % vs. the DOE),
        hypervolume in both conventions, Pareto size/composition, utopia gap, and what each
        tested sample added. `full=True` adds the HV / Pareto / utopia panels."""
        tab = self.improvement_table(through_it)
        display(Markdown(f'#### Improvement through iteration {through_it}'))
        display(tab.round(4))
        gains = self.iteration_gains(through_it)
        if len(gains):
            display(Markdown('**What each tested sample added** (vs. the best condition mean before it)'))
            display(gains.round(4))
        first, last = tab.iloc[0], tab.iloc[-1]
        print('Summary:')
        for c in Y_COLS:
            print(f"  • {Y_SHORT[c]}: best {first[f'Best {c}']:.4g} → {last[f'Best {c}']:.4g} ({last[f'Δ{c} vs DOE %']:+.2f}%)")
        print(f"  • Hypervolume (means): {first['HV (means)']:.3f} → {last['HV (means)']:.3f} ({last['ΔHV vs DOE %']:+.2f}%)"
              f"  |  replicates as points: {first['HV (reps)']:.3f} → {last['HV (reps)']:.3f}")
        print(f"  • Pareto front: {int(first['Pareto size'])} → {int(last['Pareto size'])} solutions "
              f"({int(last['Pareto from BO'])} from BO samples)  |  utopia gap {first['Utopia gap']:.3f} → {last['Utopia gap']:.3f}")
        plots.improvement(tab, full=full)
        return tab

    # ────────────────────────────── stopping rule (step by step) ─────────────
    def prb_setup(self, it, eps_frac=0.01, eps_sweep=(0.005, 0.01, 0.02, 0.05, 0.10), delta=0.05, budget=10,
                  obj_eps_frac=0.05, show=True):
        """PRB step 0 — fit the model on data through it−1 and fix the rule's parameters:
        ε (as a share of the current hypervolume), δ = δ_mod + δ_est split evenly, the planned
        budget T (δ_est^t = δ_est / (T − (it−1)) as in the paper's Appendix A.3)."""
        t_done = it - 1
        ps = self.pareto_summary(t_done)
        model, _ = fit_surrogate(self.space, ps['stats'], ps['Yvar'])        # train_Yvar = s²/n
        remaining = max(budget - t_done, 1)
        st = dict(it=it, data_through=t_done, stats=ps['stats'], Yvar=ps['Yvar'], model=model,
                  hv_now=ps['hv_means'], n_conditions=len(ps['stats']), eps_frac=eps_frac,
                  eps_sweep=sorted(set(list(eps_sweep) + [eps_frac])), eps=eps_frac * ps['hv_means'],
                  delta=delta, delta_mod=delta / 2, delta_est=delta / 2, delta_est_t=delta / 2 / remaining,
                  budget=budget, remaining=remaining, obj_eps_frac=obj_eps_frac)
        self.stopping[it] = st
        if show:
            display(Markdown(f"#### Stopping check before iteration {it} · Step 0 — model and parameters"))
            display(pd.DataFrame([{
                'Model fitted on': f"data through iteration {t_done} ({st['n_conditions']} conditions)",
                'GP noise model': 'fixed-noise GPs, train_Yvar = s²/n from the replicates (same model as the proposals)',
                'Solution judged': 'the Pareto set of the tested formulations',
                'Regret': 'largest hypervolume gain ONE more formulation could give under a joint posterior draw',
                'Current HV (means)': st['hv_now'],
                'Regret bound ε': f"{100 * eps_frac:g}% of HV = {st['eps']:.3f} HV units (sweep: {', '.join(f'{100 * f:g}%' for f in st['eps_sweep'])})",
                'Risk δ = δ_mod + δ_est': f"{delta:g} = {delta / 2:g} + {delta / 2:g}  →  λ = 1 − δ_mod = {1 - delta / 2:g}",
                'δ_est for this check': f"{st['delta_est_t']:.4g} (budget {budget} iterations, {remaining} checks remaining)",
            }]).T.rename(columns={0: 'setting'}))
            display(Markdown('**GP hyperparameters of the model being judged**'))
            display(gp_hyperparameters(model).round(4))
        return st

    def prb_draw(self, it, n_draws=2000, grid_n=40, n_show=3, show=True):
        """PRB step 1 — draw from the posterior: exact joint samples of the three latent
        objectives on a grid over the design space ∪ the tested conditions."""
        st = self.stopping[it]
        st['paths'] = prb.draw_paths(self.space, st['model'], st['stats'], n_draws=n_draws, grid_n=grid_n, seed=self.seed)
        st.update(n_draws=n_draws, grid_n=grid_n)
        if show:
            display(Markdown(f"#### Step 1 — {n_draws} joint posterior draws on a {grid_n}×{grid_n} grid ∪ {st['n_conditions']} tested conditions"))
            print(f"Each draw is one plausible version of (Y1, Y2, Y3) over the whole design space, sampled jointly "
                  f"(grid and tested points together) from the GP posterior; the arrays are ({n_draws}, {grid_n * grid_n}, 3) "
                  f"and ({n_draws}, {st['n_conditions']}, 3).")
            plots.posterior_draws(st['paths'], st['stats'], n_show=n_show)
            plots.draws_at_tested(st['paths'], st['stats'])
            chk = pd.DataFrame({'condition': [f'({a:g}, {b:g})' for a, b in st['stats'][X_COLS].values]})
            for k, c in enumerate(Y_COLS):
                chk[f'{c} posterior mean'] = st['paths']['mu_tested'][:, k]
                chk[f'{c} mean of draws'] = st['paths']['F_tested'][:, :, k].mean(0)
                chk[f'{c} sd of draws'] = st['paths']['F_tested'][:, :, k].std(0)
            display(Markdown('**Sanity check — draws reproduce the posterior at the tested conditions**'))
            display(chk.round(4))
        return st['paths']

    def prb_regret(self, it, n_front=300, examples=('median', 'max'), show=True):
        """PRB step 2 — the regret of every draw: r(f) = max_x HV(tested ∪ {x}; f) − HV(tested; f)."""
        st = self.stopping[it]
        st['regrets'] = prb.compute_regrets(self.space, st['paths'], n_front=n_front)
        r = st['regrets']['hvi_regret']
        st.update(hvi_mean=float(r.mean()), hvi_median=float(np.median(r)),
                  hv_star_mean=float(st['regrets']['hv_star'].mean()),
                  coverage=float(st['hv_now'] / st['regrets']['hv_star'].mean()))
        if show:
            display(Markdown(f"#### Step 2 — one-step hypervolume regret under each of the {st['n_draws']} draws"))
            print("For draw f: take the tested formulations at their values under f, compute the hypervolume of their Pareto\n"
                  "front, then add every grid point x alone and keep the largest hypervolume improvement. That maximum is the\n"
                  "regret of the current tested set under f — how much ONE more experiment could still gain if f were the truth.")
            order = np.argsort(r)
            pick = {'min': order[0], 'median': order[len(r) // 2], 'max': order[-1]}
            exs = [prb.example_draw(self.space, st['paths'], st['regrets'], int(pick[e])) for e in examples]
            plots.regret_examples(exs, st['stats'], st['hv_now'],
                                  title=f"Step 2 — worked examples: the {' and '.join(examples)}-regret draws")
            display(pd.DataFrame({'statistic': ['mean', 'median', '5th pct', '95th pct', 'max'],
                                  'one-step HV regret (HV units)': [r.mean(), np.median(r), np.percentile(r, 5), np.percentile(r, 95), r.max()],
                                  'as % of current HV': [100 * v / st['hv_now'] for v in (r.mean(), np.median(r), np.percentile(r, 5), np.percentile(r, 95), r.max())]}).round(4))
            gains = []
            for k in self.tested:
                prev = self.records[k - 1]['hv_means'] if (k - 1) in self.records and 'hv_means' in self.records[k - 1] \
                    else self.pareto_summary(0)['hv_means']
                if 'hv_means' in self.records[k]:
                    gains.append(self.records[k]['hv_means'] - prev)
            eps_tab = pd.DataFrame({'ε (% of HV)': [100 * f for f in st['eps_sweep']], 'ε (HV units)': [f * st['hv_now'] for f in st['eps_sweep']],
                                    'P̂(regret ≤ ε)': [(r <= f * st['hv_now']).mean() for f in st['eps_sweep']]})
            plots.regret_histogram(r, eps_tab, gains, f"Step 2 — distribution of the regret over the draws (data through iteration {st['data_through']})")
            plots.improvement_map(st['regrets']['x_best'], st['stats'], f"Step 2 — where the best additional point falls, over the {st['n_draws']} draws")
            print(f"Whole-front context: the model expects the complete Pareto front to reach HV ≈ {st['hv_star_mean']:.2f}; the tested set "
                  f"holds {100 * st['coverage']:.1f}% of it. (A continuous front is never fully covered by finitely many points, so this is "
                  f"context, not a stopping criterion.)")
        return st['regrets']

    def prb_decide(self, it, deltas=(0.01, 0.05, 0.10), budgets=(5, 10, 20), show=True):
        """PRB step 3 — the decision: sequential Clopper–Pearson test of P(regret ≤ ε) ≥ λ for
        every ε of the sweep, the test's rounds for the headline ε, and the sensitivity of the
        decision to δ and to the planned budget."""
        st = self.stopping[it]
        r = st['regrets']['hvi_regret']
        st['hvi_table'], tests = prb.prb_table(r, st['hv_now'], st['eps_sweep'], st['delta_mod'], st['delta_est_t'])
        head = tests[st['eps_frac']]
        st.update(psi=head['estimate'], decision=head['decision'].upper(), rounds=head['rounds'],
                  delta_table=prb.delta_sensitivity(r, st['hv_now'], st['eps_sweep'], deltas, st['remaining']),
                  budget_table=prb.budget_sensitivity(r, st['hv_now'], st['eps_frac'], st['delta'], budgets, st['data_through']))
        if show:
            display(Markdown(f"#### Step 3 — decision: is P(regret ≤ ε) ≥ λ = 1 − δ_mod = {1 - st['delta_mod']:g}?"))
            display(Markdown('**(a) ε sweep** — the sequential test draws in batches of 64, 96, 144, … and stops as soon as λ is '
                             'outside the Clopper–Pearson interval; STOP needs the whole interval above λ, CONTINUE the whole interval below'))
            display(st['hvi_table'].round(4))
            display(Markdown(f"**(b) The test round by round for the headline ε = {100 * st['eps_frac']:g}% of HV**"))
            display(st['rounds'].round(5))
            plots.sequential_rounds(st['rounds'], 1 - st['delta_mod'],
                                    f"Step 3 — sequential test for ε = {100 * st['eps_frac']:g}% of HV ({st['decision']})")
            display(Markdown('**(c) Sensitivity to the risk δ** (rows) for each ε (columns) — the estimate is the same, only λ and the interval width change'))
            display(st['delta_table'])
            display(Markdown(f"**(d) Sensitivity to the planned budget T** for ε = {100 * st['eps_frac']:g}% (δ_est^t = δ_est / (T − {st['data_through']}))"))
            display(st['budget_table'].round(4))
            e, p = prb.psi_curve(r, st['hv_now'])
            plots.psi_vs_eps([(f"before iteration {it} (data through {st['data_through']})", e, p)],
                             [(d, 1 - d / 2) for d in deltas], [100 * f for f in st['eps_sweep']],
                             'Step 3 — P(regret ≤ ε) as a function of ε, with λ for several δ')
        return st['hvi_table']

    def prb_objectives(self, it, show=True):
        """PRB step 4 — the paper's original single-objective rule applied to each output at its
        best tested condition (ε_k = obj_eps_frac × DOE range of that output)."""
        st = self.stopping[it]
        rng = (self.stats0[Y_COLS].max() - self.stats0[Y_COLS].min()).values
        st['obj_table'] = prb.prb_objectives(st['paths'], st['regrets'], st['stats'], st['obj_eps_frac'] * rng,
                                             st['delta_mod'], st['delta_est_t'] / 3)
        if show:
            display(Markdown(f"#### Step 4 — per-objective diagnostic (the paper's single-output rule; ε_k = {100 * st['obj_eps_frac']:g}% of the DOE range, δ_est split over 3 tests)"))
            display(st['obj_table'].round(4))
        return st['obj_table']

    def prb_conclusion(self, it):
        """PRB step 5 — the verdict in words."""
        st = self.stopping[it]
        lam = 1 - st['delta_mod']
        verdict = {'STOP': 'STOP — with probability ≥ 1 − δ under the model, no single further formulation improves the hypervolume by more than ε.',
                   'CONTINUE': 'CONTINUE — the model still expects at least one formulation that would improve the hypervolume by more than ε.',
                   'UNDECIDED': 'UNDECIDED — the estimate is too close to λ for the draws used; increase n_draws.'}[st['decision']]
        stop_from = [r['ε (% of HV)'] for _, r in st['hvi_table'].iterrows() if r['decision'] == 'STOP']
        gains = []
        for k in self.tested:
            if 'hv_means' in self.records[k]:
                prev = self.records[k - 1]['hv_means'] if (k - 1) in self.records and 'hv_means' in self.records[k - 1] \
                    else self.pareto_summary(0)['hv_means']
                gains.append(self.records[k]['hv_means'] - prev)
        obj = st.get('obj_table')
        lines = [f"CONCLUSION (stopping check before iteration {it}, data through iteration {st['data_through']}, {st['n_draws']} draws)",
                 f"  ε = {100 * st['eps_frac']:g}% of HV ({st['eps']:.3f} units), δ = {st['delta']:g}: P̂(regret ≤ ε) = {st['psi']:.3f} vs λ = {lam:g}  →  {verdict}",
                 f"  Expected best gain from one more experiment: {st['hvi_mean']:.3f} HV units ({100 * st['hvi_mean'] / st['hv_now']:.2f}% of HV)"
                 + (f"; observed gains of the tested iterations: " + ', '.join(f"{g:.3f}" for g in gains) if gains else ''),
                 f"  The rule would say STOP for ε ≥ {min(stop_from):g}% of HV" + (f" ({min(stop_from) / 100 * st['hv_now']:.2f} units)" if stop_from else '') if stop_from else "  The rule says CONTINUE for every ε in the sweep",
                 f"  Whole-front coverage: {100 * st['coverage']:.1f}% of the hypervolume the model believes attainable"]
        if obj is not None:
            done = [row['Objective'] for _, row in obj.iterrows() if row['decision'] == 'STOP']
            todo = [row['Objective'] for _, row in obj.iterrows() if row['decision'] != 'STOP']
            lines.append(f"  Per objective: optimum already found for {', '.join(done) if done else 'none'}; not yet for {', '.join(todo) if todo else 'none'}")
        lines.append("  Caveat (paper §5.3): the verdict is only as good as the model — it uses the same replicate-noise GPs that make the proposals.")
        print('\n'.join(lines))
        return st['decision']

    def stopping_check(self, it, n_draws=2000, grid_n=40, show=True, **setup_kw):
        """All PRB steps in one call (used by stopping_trajectory and for quick checks)."""
        self.prb_setup(it, show=show, **setup_kw)
        self.prb_draw(it, n_draws=n_draws, grid_n=grid_n, show=show)
        self.prb_regret(it, show=show)
        self.prb_decide(it, show=show)
        self.prb_objectives(it, show=show)
        if show:
            self.prb_conclusion(it)
        return self.stopping[it]

    def stopping_trajectory(self, through_it=None, eps_sweep=(0.01, 0.02, 0.05, 0.10), deltas=(0.01, 0.05, 0.10), **kw):
        """How the rule evolved: the same check as it would have been made before each request
        (model on data through 0, 1, …). Missing stages are computed with the settings of the
        latest check (n_draws, grid, δ, budget); existing ones are reused."""
        through_it = through_it or max(self.records)
        ref = self.stopping[max(self.stopping)] if self.stopping else None
        for it in range(1, through_it + 1):
            st = self.stopping.get(it)
            if st is None or 'hvi_table' not in st or not all(any(np.isclose(st['hvi_table']['ε (% of HV)'], 100 * f)) for f in eps_sweep):
                base = dict(eps_frac=ref['eps_frac'], delta=ref['delta'], budget=ref['budget'], obj_eps_frac=ref['obj_eps_frac'],
                            n_draws=ref['n_draws'], grid_n=ref['grid_n']) if ref else {}
                base.update(kw); base['eps_sweep'] = sorted(set(list(eps_sweep) + [base.get('eps_frac', 0.01)]))
                self.stopping_check(it, show=False, **base)
        rows, curves, hists = [], [], []
        for it in range(1, through_it + 1):
            st = self.stopping[it]; r = st['regrets']['hvi_regret']
            for _, row in st['hvi_table'].iterrows():
                if any(np.isclose(row['ε (% of HV)'], 100 * f) for f in eps_sweep):
                    rows.append({'Check before iteration': it, 'Data through': it - 1, 'Conditions': st['n_conditions'],
                                 'HV (means)': st['hv_now'], 'ε (% of HV)': row['ε (% of HV)'],
                                 'P̂(regret ≤ ε)': row['P̂(regret ≤ ε)'], 'decision': row['decision'],
                                 'E[max HVI]': st['hvi_mean'], 'median regret': st['hvi_median'],
                                 'E[HV*]': st['hv_star_mean'], 'front coverage': st['coverage']})
            e, p = prb.psi_curve(r, st['hv_now'])
            curves.append((f'before iteration {it} (data through {it - 1})', e, p)); hists.append((f'before iteration {it}', r))
        traj = pd.DataFrame(rows)
        display(Markdown('**P̂(one-step HV regret ≤ ε) before each request**'))
        display(traj.pivot(index='Check before iteration', columns='ε (% of HV)', values='P̂(regret ≤ ε)').round(3)
                .rename(columns=lambda c: f'ε = {c:g}% HV'))
        display(Markdown('**Decision before each request**'))
        display(traj.pivot(index='Check before iteration', columns='ε (% of HV)', values='decision')
                .rename(columns=lambda c: f'ε = {c:g}% HV'))
        display(Markdown('**Expected gain of one more experiment and front coverage, stage by stage**'))
        display(traj.drop_duplicates('Check before iteration')
                [['Check before iteration', 'Data through', 'Conditions', 'HV (means)', 'E[max HVI]', 'median regret', 'E[HV*]', 'front coverage']]
                .set_index('Check before iteration').round(4))
        lam = 1 - self.stopping[through_it]['delta_mod']
        plots.stopping_trajectory(traj, lam)
        plots.regret_by_stage(hists)
        plots.psi_vs_eps(curves, [(d, 1 - d / 2) for d in deltas], [100 * f for f in eps_sweep],
                         'How the rule evolved: P(regret ≤ ε) vs ε at each stage (curves move up-left as the front is learned)')
        return traj

    # ────────────────────────────── summaries ────────────────────────────────
    def improvement_table(self, through_it):
        lo = self.stats0[Y_COLS].min().values
        span = (self.stats0[Y_COLS].max() - self.stats0[Y_COLS].min()).values
        rows = []
        for t in range(through_it + 1):
            ps = self.pareto_summary(t); st = ps['stats']
            r = {'Iteration': t, 'Conditions': len(st)}
            for c in Y_COLS:
                r[f'Best {c}'] = st[c].min()
            r['HV (means)'] = ps['hv_means']; r['HV (reps)'] = ps['hv_reps']
            r['Pareto size'] = ps['n_pareto']; r['Pareto from BO'] = int((st.loc[ps['pm'], 'Iteration'] > 0).sum())
            r['Utopia gap'] = float(np.sqrt((((st[Y_COLS].values - lo) / span) ** 2).sum(axis=1)).min())
            rows.append(r)
        tab = pd.DataFrame(rows)
        for c in Y_COLS:
            tab[f'Δ{c} vs DOE %'] = 100 * (tab[f'Best {c}'] / tab.loc[0, f'Best {c}'] - 1)
        tab['ΔHV vs DOE %'] = 100 * (tab['HV (means)'] / tab.loc[0, 'HV (means)'] - 1)
        return tab

    def iteration_gains(self, through_it):
        rows = []
        for t in range(1, through_it + 1):
            ps = self.pareto_summary(t); st = ps['stats']
            before, new = st[st.Iteration < t], st[st.Iteration == t]
            for idx, s in new.iterrows():
                r = {'Iteration': t, 'X1': s.X1, 'X2': s.X2, 'On Pareto front': bool(ps['pm'][idx])}
                for c in Y_COLS:
                    r[f'{c} mean'] = s[c]; r[f'{c} best before'] = before[c].min()
                    r[f'{c} improved'] = bool(s[c] < before[c].min())
                rows.append(r)
        return pd.DataFrame(rows)

    def table(self):
        """One row per iteration: requested vs. fabricated, sent / v4 / observed means, HV after."""
        rows = []
        for it in sorted(self.records):
            rec = self.records[it]
            r = {'Iteration': it, 'Requested X1': rec['x_proposed'][0, 0], 'Requested X2': rec['x_proposed'][0, 1],
                 'Fabricated': f"{rec['x_fab'][0, 0]:g} / {rec['x_fab'][0, 1]:g}", 'qLogEHVI': rec['ehvi'],
                 'Status': 'tested' if 'observed' in rec else 'pending'}
            for k, c in enumerate(Y_COLS):
                r[f'{c} sent µ'] = rec['sent'].loc[k, 'mean'] if rec.get('sent') is not None else np.nan
                r[f'{c} v4 µ'] = rec['pred'].loc[k, 'mean']
                r[f'{c} observed'] = rec['observed'][c].mean() if 'observed' in rec else np.nan
            r['HV after (means)'] = rec.get('hv_means', np.nan); r['HV after (reps)'] = rec.get('hv_reps', np.nan)
            r['Origin'] = rec['origin']
            rows.append(r)
        return pd.DataFrame(rows)

    def performance_table(self, compact=False):
        """Requested vs. provided, one row per iteration and objective: the composition we
        requested, the prediction (and expected range) that was SENT to UIC, what UIC measured,
        the difference, and whether the measurements fell inside the range we gave them."""
        rows = []
        for it in self.tested:
            rec = self.records[it]; obs = rec['observed']; sent = rec.get('sent'); pred = rec['pred']
            for k, c in enumerate(Y_COLS):
                o = obs[c].values
                if sent is not None:
                    m, lo, hi = float(sent.loc[k, 'mean']), float(sent.loc[k, 'low']), float(sent.loc[k, 'high'])
                    kind = sent.attrs.get('kind', 'sent range')
                else:
                    m, lo, hi = float(pred.loc[k, 'mean']), float(pred.loc[k, 'pi_low']), float(pred.loc[k, 'pi_high'])
                    kind = 'v4 95% prediction interval (nothing was sent)'
                inside = int(((o >= lo) & (o <= hi)).sum())
                verdict = ('✓ all replicates inside the range' if inside == len(o) else
                           ('~ mean inside, scatter outside' if lo <= o.mean() <= hi else '✗ mean outside the sent range'))
                rows.append({'Iteration': it,
                             'Requested (X1, X2)': f"{rec['x_proposed'][0, 0]:.2f}, {rec['x_proposed'][0, 1]:.2f}",
                             'Fabricated (carbon black / regolith wt%)': f"{rec['x_fab'][0, 0]:g} / {rec['x_fab'][0, 1]:g}",
                             'Objective': Y_NAMES[c], 'Predicted (sent)': m, 'Expected range (sent)': f'{lo:.4g} – {hi:.4g}',
                             'Range type': kind, 'Measured replicates': ', '.join(f'{v:g}' for v in o), 'Measured mean': o.mean(),
                             'Measured − predicted': o.mean() - m, 'Error (% of predicted)': 100 * abs(o.mean() - m) / m,
                             'Replicates inside range': f'{inside} of {len(o)}', 'Verdict': verdict,
                             'v4 model mean': float(pred.loc[k, 'mean']), 'v4 95% PI': f"{pred.loc[k, 'pi_low']:.4g} – {pred.loc[k, 'pi_high']:.4g}"})
        df = pd.DataFrame(rows)
        if compact:
            df = df[['Iteration', 'Fabricated (carbon black / regolith wt%)', 'Objective', 'Predicted (sent)', 'Expected range (sent)',
                     'Measured replicates', 'Measured mean', 'Error (% of predicted)', 'Replicates inside range', 'Verdict']].copy()
            df['Predicted (sent)'] = df['Predicted (sent)'].map(lambda v: f'{v:.4g}')
            df['Measured mean'] = df['Measured mean'].map(lambda v: f'{v:.4g}')
            df['Error (% of predicted)'] = df['Error (% of predicted)'].map(lambda v: f'{v:.1f}%')
            df = df.rename(columns={'Fabricated (carbon black / regolith wt%)': 'Formulation (CB/regolith wt%)'})
        return df

    def performance_figure(self, path=None):
        """Predicted-as-sent vs. measured, per objective and iteration (shareable)."""
        rows = []
        for it in self.tested:
            rec = self.records[it]; sent = rec.get('sent'); pred = rec['pred']
            for k, c in enumerate(Y_COLS):
                src = (float(sent.loc[k, 'mean']), float(sent.loc[k, 'low']), float(sent.loc[k, 'high'])) if sent is not None \
                    else (float(pred.loc[k, 'mean']), float(pred.loc[k, 'pi_low']), float(pred.loc[k, 'pi_high']))
                rows.append(dict(it=it, c=c, mean=src[0], low=src[1], high=src[2], obs=rec['observed'][c].values,
                                 fab=f"{rec['x_fab'][0, 0]:g}/{rec['x_fab'][0, 1]:g}"))
        if path: os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
        plots.model_performance(rows, path=path)

    def stopping_summary_figure(self, traj, path=None):
        """The one figure to share about stopping (see plots.stopping_summary)."""
        gains = []
        for k in self.tested:
            if 'hv_means' in self.records[k]:
                prev = self.records[k - 1]['hv_means'] if (k - 1) in self.records and 'hv_means' in self.records[k - 1] \
                    else self.pareto_summary(0)['hv_means']
                gains.append((k, 100 * (self.records[k]['hv_means'] - prev) / prev))
        lam = 1 - self.stopping[max(self.stopping)]['delta_mod']
        if path: os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
        plots.stopping_summary(traj, gains, lam, path=path)

    def prediction_accuracy(self, plot=True):
        rows = []
        for it in self.tested:
            rec = self.records[it]
            for k, c in enumerate(Y_COLS):
                o = rec['observed'][c].values; p = rec['pred'].loc[k]
                r = {'Iteration': it, 'Objective': Y_SHORT[c], 'Obs mean': o.mean(), 'v4 µ': p['mean'],
                     'Δ obs − v4 µ': o.mean() - p['mean'], '|Δ| / σ_pred (v4)': abs(o.mean() - p['mean']) / p['sd_pred'],
                     'reps in v4 PI': int(((o >= p['pi_low']) & (o <= p['pi_high'])).sum())}
                if rec.get('sent') is not None:
                    s = rec['sent'].loc[k]
                    r['Sent µ'] = s['mean']; r['Δ obs − sent µ'] = o.mean() - s['mean']
                    r['reps in sent range'] = int(((o >= s['low']) & (o <= s['high'])).sum())
                rows.append(r)
        acc = pd.DataFrame(rows)
        if plot:
            display(acc.round(4)); plots.prediction_accuracy(acc)
        return acc

    def response_surfaces(self, it=None):
        """GP surfaces of the model that proposed iteration `it` (default: the pending one),
        with the tested samples and that proposal marked."""
        it = it or max(self.records)
        rec = self.records[it]
        plots.response_surfaces(self.space, rec['model'], rec['stats'], proposed=rec['x_proposed'][0],
                                title=f'v4 GP response surfaces (data through iteration {it - 1}) — campaign points highlighted')

    def export(self, next_it=None):
        """Write outputs/v4: bo_iterations.xlsx, ablation_data_by_iteration.xlsx and the
        one-sheet recommendation workbook for the pending iteration."""
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        next_it = next_it or max(self.records)
        paths, R = [], self.records

        preds = pd.concat([R[it]['pred'].assign(iteration=it) for it in sorted(R)], ignore_index=True)
        sent = pd.concat([R[it]['sent'].assign(iteration=it, source=R[it]['sent_source'])
                          for it in sorted(R) if R[it].get('sent') is not None], ignore_index=True)
        observed = pd.concat([R[it]['observed'].assign(iteration=it, replicate=range(1, len(R[it]['observed']) + 1),
                                                       source=R[it]['source']) for it in self.tested], ignore_index=True)
        info = pd.DataFrame({'field': ['version', 'model', 'acquisition', 'source code', 'intervals', 'hypervolume', 'data file'],
                             'value': ['v4', 'fixed-noise SingleTaskGP per objective (train_Yvar = s²/n from replicates)',
                                       f'qLogEHVI, optimize_acqf(num_restarts=20, raw_samples=256), seed {self.seed}',
                                       'notebooks/mobo_analysis_v4.ipynb + notebooks/utils/',
                                       'ci = latent-mean 95% CI; pi = 95% prediction interval for one replicate (latent + pooled s̄²)',
                                       'reference point = initial-DOE nadir + 10% slack; HV (means) on condition means, '
                                       'HV (reps) with campaign replicates as points', self.data_file]})
        p = os.path.join(self.out_dir, 'bo_iterations.xlsx')
        with pd.ExcelWriter(p) as xl:
            self.table().to_excel(xl, sheet_name='summary', index=False)
            preds[['iteration'] + [c for c in preds.columns if c != 'iteration']].to_excel(xl, sheet_name='predictions_v4', index=False)
            sent[['iteration'] + [c for c in sent.columns if c != 'iteration']].to_excel(xl, sheet_name='predictions_sent', index=False)
            observed[['iteration', 'replicate'] + X_COLS + Y_COLS + ['source']].to_excel(xl, sheet_name='observed', index=False)
            self.prediction_accuracy(plot=False).to_excel(xl, sheet_name='accuracy', index=False)
            self.improvement_table(max(self.tested)).to_excel(xl, sheet_name='improvement', index=False)
            info.to_excel(xl, sheet_name='info', index=False)
        paths.append(p)

        df_all = self.replicates_through(max(self.tested))
        df_all.insert(0, 'Iteration', [self.iteration_of(r.X1, r.X2) for r in df_all.itertuples()])
        p = os.path.join(self.out_dir, 'ablation_data_by_iteration.xlsx')
        df_all.rename(columns={**X_NAMES, **Y_NAMES}).to_excel(p, index=False)
        fills = {1: 'FFF2CC', 2: 'D9EAD3', 3: 'CFE2F3', 4: 'F4CCCC', 5: 'EAD1DC'}
        wb = load_workbook(p); ws = wb.active
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            if row[0].value in fills:
                for c in row:
                    c.fill = PatternFill('solid', fgColor=fills[row[0].value])
        for col_cells in ws.columns:
            ws.column_dimensions[col_cells[0].column_letter].width = max(len(str(c.value)) for c in col_cells if c.value is not None) + 2
        wb.save(p); paths.append(p)

        share = os.path.join(self.out_dir, 'share'); os.makedirs(share, exist_ok=True)
        perf = self.performance_table()
        p = os.path.join(share, 'model_performance.xlsx')
        with pd.ExcelWriter(p) as xl:
            perf.to_excel(xl, sheet_name='requested_vs_provided', index=False)
            self.performance_table(compact=True).to_excel(xl, sheet_name='compact', index=False)
        paths.append(p)
        p = os.path.join(share, 'model_performance_table.png')
        plots.table_png(self.performance_table(compact=True), p, 'Requested vs. provided — model predictions against UIC measurements',
                        col_widths=[0.05, 0.13, 0.15, 0.08, 0.11, 0.13, 0.08, 0.09, 0.10, 0.17])
        paths.append(p)

        rec = R[next_it]; pr = rec['pred']
        x1f, x2f = float(rec['x_fab'][0, 0]), float(rec['x_fab'][0, 1])
        rows = [(f'Recommended sample (iteration {next_it})', None, None), ('Component', 'wt%', None),
                ('Carbon black', x1f, None), ('Fe-rich regolith', x2f, None), ('I-369', I369_WT, None),
                ('Resin', round(100 - x1f - x2f - I369_WT, 1), None), (None, None, None), ('Predicted values', None, None),
                ('Objective', 'Predicted mean', 'Expected range per replicate (95%)'),
                (Y_NAMES['Y1'], round(float(pr.loc[0, 'mean']), 2), f"{pr.loc[0, 'pi_low']:.1f} – {pr.loc[0, 'pi_high']:.1f}"),
                (Y_NAMES['Y2'], round(float(pr.loc[1, 'mean']), 1), f"{pr.loc[1, 'pi_low']:.1f} – {pr.loc[1, 'pi_high']:.1f}"),
                (Y_NAMES['Y3'], round(float(pr.loc[2, 'mean']), 4), f"{pr.loc[2, 'pi_low']:.4f} – {pr.loc[2, 'pi_high']:.4f}")]
        wb = Workbook(); ws = wb.active; ws.title = 'recommendation'
        for r in rows:
            ws.append(list(r))
        for coord in ('A1', 'A2', 'B2', 'A8', 'A9', 'B9', 'C9'):
            ws[coord].font = Font(bold=True)
        for col, w in zip('ABC', (28, 15, 34)):
            ws.column_dimensions[col].width = w
        p = os.path.join(self.out_dir, f'iteration{next_it}_recommended_sample.xlsx'); wb.save(p); paths.append(p)

        display(Markdown(f'**Recommendation for UIC (iteration {next_it})**'))
        display(pd.DataFrame(rows[2:6], columns=['Component', 'wt%', '_']).iloc[:, :2].set_index('Component').T)
        display(pd.DataFrame(rows[9:12], columns=list(rows[8])).set_index('Objective'))
        for p in paths:
            print('wrote', p)
        return paths

    def summary(self):
        """Plain-text status block: requested → received per iteration, the pending request,
        and the improvement vs. the initial DOE."""
        R, tested, nxt = self.records, self.tested, max(self.records)
        tab = self.improvement_table(max(tested)); first, last = tab.iloc[0], tab.iloc[-1]
        lines = [f'CAMPAIGN STATUS — {len(tested)} tested iteration(s), iteration {nxt} '
                 f"{'pending' if nxt not in tested else 'tested'}", '', 'Requested → received:']
        for it in tested:
            rec, o = R[it], R[it]['observed']
            lines.append(f"  iter {it}: requested ({rec['x_proposed'][0, 0]:.2f}, {rec['x_proposed'][0, 1]:.2f}) → fabricated "
                         f"{rec['x_fab'][0, 0]:g}/{rec['x_fab'][0, 1]:g}; observed means "
                         + ', '.join(f"{c} {o[c].mean():.4g} (v4 µ {rec['pred'].loc[k, 'mean']:.4g})" for k, c in enumerate(Y_COLS))
                         + f"; HV after {rec['hv_means']:.3f}")
        if nxt not in tested:
            rec = R[nxt]
            lines += ['', f"Next request (iteration {nxt}): X1 = {rec['x_proposed'][0, 0]:.3f}, X2 = {rec['x_proposed'][0, 1]:.3f} "
                          f"→ {rec['x_fab'][0, 0]:g} / {rec['x_fab'][0, 1]:g}  (qLogEHVI {rec['ehvi']:.4f})",
                      '  predicted: ' + ' | '.join(f"{c} {rec['pred'].loc[k, 'mean']:.4g} PI [{rec['pred'].loc[k, 'pi_low']:.4g}, "
                                                   f"{rec['pred'].loc[k, 'pi_high']:.4g}]" for k, c in enumerate(Y_COLS))]
        if nxt in self.stopping:
            st = self.stopping[nxt]
            lines += ['', f"Stopping rule before iteration {nxt} (Wilson 2024, one-step HV regret; ε = {100 * st['eps_frac']:g}% of HV, "
                          f"δ = {st['delta']:g}): P̂(regret ≤ ε) = {st['psi']:.3f} vs λ = {1 - st['delta_mod']:g} → {st['decision']}",
                      f"  expected best one-step gain {st['hvi_mean']:.3f} HV units ({100 * st['hvi_mean'] / st['hv_now']:.2f}% of HV); "
                      f"front coverage {100 * st['coverage']:.1f}%; "
                      + '; '.join(f"ε={r['ε (% of HV)']:g}%: {r['decision']}" for _, r in st['hvi_table'].iterrows())]
        lines += ['', 'Improvement vs. the initial DOE (best condition mean):']
        for c in Y_COLS:
            lines.append(f"  {Y_SHORT[c]}: {first[f'Best {c}']:.4g} → {last[f'Best {c}']:.4g} ({last[f'Δ{c} vs DOE %']:+.2f}%)")
        lines += [f"  Hypervolume (means): {first['HV (means)']:.3f} → {last['HV (means)']:.3f} ({last['ΔHV vs DOE %']:+.2f}%); "
                  f"replicates as points: {first['HV (reps)']:.3f} → {last['HV (reps)']:.3f}",
                  f"  Pareto front: {int(first['Pareto size'])} → {int(last['Pareto size'])} solutions, "
                  f"{int(last['Pareto from BO'])} contributed by BO samples"]
        print('\n'.join(lines))

    # ────────────────────────────── helpers ──────────────────────────────────
    @staticmethod
    def _show_prediction(pred, title):
        t = pred.copy()
        t['95% CI (latent mean)'] = [f'[{a:.4f}, {b:.4f}]' for a, b in zip(t.ci_low, t.ci_high)]
        t['95% PI (one replicate)'] = [f'[{a:.4f}, {b:.4f}]' for a, b in zip(t.pi_low, t.pi_high)]
        display(Markdown(f'**{title}**'))
        display(t[['Objective', 'mean', 'sd_latent', '95% CI (latent mean)', 'sd_pred', '95% PI (one replicate)']].round(4))
